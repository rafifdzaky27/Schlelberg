from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parents[1]
CHAPTERS = sorted((ROOT / "work" / "manuscript").glob("[0-9][0-9]-*.md"))
OUTPUT = ROOT / "outputs"
ART = ROOT / "work" / "cover" / "schlelberg-art-v2.png"


def add_runs(paragraph, source):
    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", source):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        else:
            paragraph.add_run(piece)


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def make_interior():
    doc = Document()
    first = doc.sections[0]
    for sec in [first]:
        sec.page_width = Inches(6)
        sec.page_height = Inches(9)
        sec.top_margin = Inches(0.88)
        sec.bottom_margin = Inches(0.73)
        sec.left_margin = Inches(0.86)
        sec.right_margin = Inches(0.72)
        sec.header_distance = Inches(0.38)
        sec.footer_distance = Inches(0.39)

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(20, 24, 28)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(14.2)
    normal.paragraph_format.first_line_indent = Inches(0.20)
    normal.paragraph_format.widow_control = True

    for stylename in ["Title", "Subtitle", "Heading 1"]:
        style = doc.styles[stylename]
        style.font.name = "Baskerville"
        style.font.color.rgb = RGBColor(20, 24, 28)
        ppr = style._element.pPr
        if ppr is not None:
            for border in ppr.findall(qn("w:pBdr")):
                ppr.remove(border)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Inches(2.5)
    title.paragraph_format.space_after = Pt(20)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("SCHLELBERG")
    title_run.font.name = "Baskerville"
    title_run.font.size = Pt(29)

    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Inches(1.0)
    sub_run = sub.add_run("A NOVEL")
    sub_run.font.name = "Baskerville"
    sub_run.font.size = Pt(11)
    sub_run.italic = False

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.first_line_indent = Inches(0)
    run = author.add_run("RAFIF")
    run.font.name = "Baskerville"
    run.font.size = Pt(14)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.page_width = first.page_width
    body_section.page_height = first.page_height
    body_section.top_margin = first.top_margin
    body_section.bottom_margin = first.bottom_margin
    body_section.left_margin = first.left_margin
    body_section.right_margin = first.right_margin
    body_section.footer_distance = first.footer_distance
    body_section.footer.is_linked_to_previous = False
    set_page_number(body_section.footer.paragraphs[0])
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), "1")
    body_section._sectPr.append(pg_num)

    for index, path in enumerate(CHAPTERS):
        lines = path.read_text(encoding="utf-8-sig").splitlines()
        heading = next((line[3:] for line in lines if line.startswith("## ")), "")
        number, chapter_title = heading.split(" — ", 1)
        if index:
            doc.add_page_break()
        num = doc.add_paragraph()
        num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        num.paragraph_format.first_line_indent = Inches(0)
        num.paragraph_format.space_before = Inches(0.66)
        num.paragraph_format.space_after = Pt(10)
        num.paragraph_format.keep_with_next = True
        r = num.add_run(number.upper())
        r.font.name = "Georgia"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(103, 103, 101)

        ttl = doc.add_paragraph(style="Heading 1")
        ttl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ttl.paragraph_format.first_line_indent = Inches(0)
        ttl.paragraph_format.space_after = Pt(24)
        ttl.paragraph_format.keep_with_next = True
        tr = ttl.add_run(chapter_title)
        tr.font.name = "Baskerville"
        tr.font.size = Pt(19)
        tr.font.bold = False

        prose = "\n".join(line for line in lines if not line.startswith("# ") and not line.startswith("## ")).strip()
        blocks = re.split(r"\n\s*\n", prose)
        next_opening = True
        for block in blocks:
            block = block.strip().replace("\n", " ")
            if not block:
                continue
            if block == "***":
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_before = Pt(11)
                p.paragraph_format.space_after = Pt(11)
                p.add_run("*   *   *")
                next_opening = True
                continue
            p = doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if next_opening:
                p.paragraph_format.first_line_indent = Inches(0)
                next_opening = False
            add_runs(p, block)

    doc.core_properties.title = "Schlelberg"
    doc.core_properties.author = "Rafif"
    doc.core_properties.subject = "A novel"
    OUTPUT.mkdir(exist_ok=True)
    doc.save(OUTPUT / "Schlelberg - novel.docx")


def wrap_text(c, text, x, y, max_width, size=11.5, leading=17):
    c.setFont("Georgia", size)
    words = text.split()
    line = ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if pdfmetrics.stringWidth(candidate, "Georgia", size) > max_width and line:
            c.drawString(x, y, line)
            y -= leading
            line = word
        else:
            line = candidate
    if line:
        c.drawString(x, y, line)
        y -= leading
    return y


def make_jacket():
    inch = 72
    back = 6.25 * inch
    spine = 0.75 * inch
    front = 6.25 * inch
    height = 9.25 * inch
    width = back + spine + front
    c = canvas.Canvas(str(OUTPUT / "Schlelberg - hardcover jacket concept.pdf"), pagesize=(width, height))
    c.setTitle("Schlelberg hardcover jacket concept")
    c.setAuthor("Rafif")
    pdfmetrics.registerFont(TTFont("Baskerville", "C:/Windows/Fonts/BASKVILL.TTF"))
    pdfmetrics.registerFont(TTFont("Georgia", "C:/Windows/Fonts/georgia.ttf"))

    # Back and spine continue the darkest color in the illustrated sky.
    c.setFillColorRGB(0.075, 0.13, 0.20)
    c.rect(0, 0, width, height, fill=1, stroke=0)
    front_x = back + spine
    c.saveState()
    clip = c.beginPath()
    clip.rect(front_x, 0, front, height)
    c.clipPath(clip, stroke=0)
    image = ImageReader(str(ART))
    iw, ih = image.getSize()
    scale = max(front / iw, height / ih)
    draw_w, draw_h = iw * scale, ih * scale
    c.drawImage(image, front_x + (front - draw_w) / 2, (height - draw_h) / 2, draw_w, draw_h)
    c.restoreState()

    c.saveState()
    c.setFillColorRGB(0.025, 0.055, 0.11)
    c.setFillAlpha(0.52)
    c.rect(front_x, height - 168, front, 168, fill=1, stroke=0)
    c.restoreState()
    c.setFillColorRGB(0.96, 0.94, 0.88)
    c.setFont("Baskerville", 43)
    c.drawCentredString(front_x + front / 2, height - 81, "SCHLELBERG")
    c.setFont("Georgia", 10)
    c.drawCentredString(front_x + front / 2, height - 105, "A NOVEL")
    c.setFont("Baskerville", 18)
    c.drawCentredString(front_x + front / 2, 43, "RAFIF")

    c.setStrokeColorRGB(0.60, 0.53, 0.37)
    c.setLineWidth(0.6)
    c.line(back, 0, back, height)
    c.line(back + spine, 0, back + spine, height)
    c.saveState()
    c.translate(back + spine / 2, 55)
    c.rotate(90)
    c.setFillColorRGB(0.96, 0.94, 0.88)
    c.setFont("Baskerville", 17)
    c.drawString(0, -6, "SCHLELBERG")
    c.setFont("Georgia", 8)
    c.drawRightString(height - 110, -4, "RAFIF")
    c.restoreState()

    x = 48
    c.setFillColorRGB(0.89, 0.79, 0.59)
    c.setFont("Baskerville", 21)
    c.drawString(x, height - 80, "Two shores. One failing door.")
    c.setFillColorRGB(0.94, 0.93, 0.87)
    y = height - 123
    paragraphs = [
        "On the ninth day of Aras's occupation, Tobious opens a door buried beneath his village and glimpses a wounded version of himself in a city twenty-six hundred years away.",
        "The future offers the strength he was denied: a shield, a legendary blade, and a chance to bring Mariya home. But the gate is built on lives its keepers have chosen not to count. As the Hollow Ones gather in the space between centuries, every rescue shifts danger onto someone unseen.",
        "Mariya is keeping their village alive by other means: grain, water, names, and the stubborn belief that no one should vanish from the record. To save both shores, they must decide what can be kept and what love cannot.",
    ]
    for paragraph in paragraphs:
        y = wrap_text(c, paragraph, x, y, back - 2 * x, size=11.5, leading=18)
        y -= 16
    c.setStrokeColorRGB(0.59, 0.50, 0.34)
    c.line(x, 76, back - x, 76)
    c.setFillColorRGB(0.89, 0.79, 0.59)
    c.setFont("Georgia", 8)
    c.drawString(x, 55, "AN EPIC OF TWO ERAS")
    c.showPage()
    c.save()


if __name__ == "__main__":
    assert len(CHAPTERS) == 32, len(CHAPTERS)
    make_interior()
    make_jacket()
